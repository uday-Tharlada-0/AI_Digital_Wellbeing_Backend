import io
from datetime import date, timedelta
from flask import Blueprint, jsonify, request, send_file
from sqlalchemy import func

from extensions import db
from models import Application, UsageRecord, HourlyActivity, FocusSession
from utils import fmt_minutes, productivity_score, date_range, TIME_BUCKETS

analytics_bp = Blueprint("analytics", __name__, url_prefix="/api/analytics")


def _day_totals(the_date):
    rows = (
        db.session.query(UsageRecord.minutes, Application.is_productive)
        .join(Application, Application.id == UsageRecord.application_id)
        .filter(UsageRecord.date == the_date)
        .all()
    )
    total = sum(r[0] for r in rows)
    productive = sum(r[0] for r in rows if r[1])
    return total, productive


@analytics_bp.route("/daily-trends", methods=["GET"])
def daily_trends():
    days = int(request.args.get("days", 7))
    days = min(max(days, 1), 90)
    out = []
    for d in date_range(days):
        total, productive = _day_totals(d)
        out.append(
            {
                "date": d.isoformat(),
                "label": d.strftime("%b %d"),
                "total_hours": round(total / 60, 2),
                "productive_hours": round(productive / 60, 2),
                "productivity_score": productivity_score(productive, total),
            }
        )
    return jsonify(out)


@analytics_bp.route("/time-of-day", methods=["GET"])
def time_of_day():
    days = int(request.args.get("days", 7))
    start = date_range(days)[0]

    rows = (
        db.session.query(HourlyActivity.hour, func.sum(HourlyActivity.minutes))
        .filter(HourlyActivity.date >= start)
        .group_by(HourlyActivity.hour)
        .all()
    )
    hour_totals = {h: m for h, m in rows}

    bucket_totals = {label: 0.0 for label in TIME_BUCKETS}
    for hour, minutes in hour_totals.items():
        for label, hours in TIME_BUCKETS.items():
            if hour in hours:
                bucket_totals[label] += minutes
                break

    return jsonify(
        [
            {"bucket": label, "hours": round(minutes / 60, 2)}
            for label, minutes in bucket_totals.items()
        ]
    )


@analytics_bp.route("/week-over-week", methods=["GET"])
def week_over_week():
    weeks = int(request.args.get("weeks", 4))
    out = []
    today = date.today()
    for w in range(weeks - 1, -1, -1):
        week_end = today - timedelta(days=7 * w)
        week_start = week_end - timedelta(days=6)
        total = (
            db.session.query(func.sum(UsageRecord.minutes))
            .filter(UsageRecord.date >= week_start, UsageRecord.date <= week_end)
            .scalar()
            or 0
        )
        out.append({"label": f"W{weeks - w}", "hours": round(total / 60, 2)})
    return jsonify(out)


@analytics_bp.route("/insights", methods=["GET"])
def insights():
    today = date.today()
    week_start = date_range(7)[0]

    total_week, productive_week = 0, 0
    for d in date_range(7):
        t, p = _day_totals(d)
        total_week += t
        productive_week += p
    score_week = productivity_score(productive_week, total_week)

    # peak hour
    peak_row = (
        db.session.query(HourlyActivity.hour, func.sum(HourlyActivity.minutes))
        .filter(HourlyActivity.date >= week_start)
        .group_by(HourlyActivity.hour)
        .order_by(func.sum(HourlyActivity.minutes).desc())
        .first()
    )

    # biggest week-over-week category increase
    this_week = dict(
        db.session.query(Application.category, func.sum(UsageRecord.minutes))
        .join(UsageRecord, UsageRecord.application_id == Application.id)
        .filter(UsageRecord.date >= week_start)
        .group_by(Application.category)
        .all()
    )
    prev_start = week_start - timedelta(days=7)
    prev_end = week_start - timedelta(days=1)
    prev_week = dict(
        db.session.query(Application.category, func.sum(UsageRecord.minutes))
        .join(UsageRecord, UsageRecord.application_id == Application.id)
        .filter(UsageRecord.date >= prev_start, UsageRecord.date <= prev_end)
        .group_by(Application.category)
        .all()
    )

    biggest_increase, biggest_pct = None, 0
    for cat, mins in this_week.items():
        prev = prev_week.get(cat, 0)
        if prev > 0:
            pct = (mins - prev) / prev * 100
            if pct > biggest_pct:
                biggest_pct, biggest_increase = pct, cat

    # focus completion rate this week
    sessions = FocusSession.query.filter(FocusSession.date >= week_start).all()
    completed = [s for s in sessions if s.completed]
    completion_rate = round(len(completed) / len(sessions) * 100) if sessions else None

    out = []
    if score_week >= 60:
        out.append(
            {
                "icon": "🎯",
                "text": f"Excellent productivity: {score_week}% of your screen time this week was productive work.",
            }
        )
    elif total_week > 0:
        out.append(
            {
                "icon": "⚖️",
                "text": f"Your productivity score this week is {score_week}% — try scheduling a daily focus block to raise it.",
            }
        )

    if peak_row and peak_row[0] is not None:
        hour = peak_row[0]
        label = f"{hour % 12 or 12} {'AM' if hour < 12 else 'PM'}"
        out.append({"icon": "⏰", "text": f"Your most active time is around {label} — plan deep work then."})

    if biggest_increase:
        out.append(
            {
                "icon": "📈",
                "text": f"{biggest_increase} usage is up {round(biggest_pct)}% week-over-week — consider a limit if it's leisure time.",
            }
        )

    if completion_rate is not None:
        out.append(
            {
                "icon": "🍅",
                "text": f"Focus session completion rate is {completion_rate}% this week.",
            }
        )

    if not out:
        out.append({"icon": "👋", "text": "Not enough data yet — keep tracking to unlock personalized insights."})

    return jsonify(out)


@analytics_bp.route("/export", methods=["GET"])
def export_excel():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    days = int(request.args.get("days", 30))
    wb = Workbook()

    # Summary sheet
    ws = wb.active
    ws.title = "Summary"
    header_fill = PatternFill("solid", fgColor="4A6CF7")
    ws.append(["Date", "Total Hours", "Productive Hours", "Productivity Score"])
    for cell in ws[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill

    for d in date_range(days):
        total, productive = _day_totals(d)
        ws.append(
            [
                d.isoformat(),
                round(total / 60, 2),
                round(productive / 60, 2),
                productivity_score(productive, total),
            ]
        )

    # Apps sheet
    ws2 = wb.create_sheet("Applications")
    ws2.append(["Application", "Category", "Productive", f"Total Minutes ({days}d)"])
    for cell in ws2[1]:
        cell.font = Font(bold=True)

    week_start = date_range(days)[0]
    for app in Application.query.all():
        total = (
            db.session.query(func.sum(UsageRecord.minutes))
            .filter(UsageRecord.application_id == app.id, UsageRecord.date >= week_start)
            .scalar()
            or 0
        )
        ws2.append([app.name, app.category, "Yes" if app.is_productive else "No", round(total, 1)])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(
        buf,
        as_attachment=True,
        download_name=f"scolect_report_{date.today().isoformat()}.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
